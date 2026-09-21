# -*- coding: utf-8 -*-
"""Dựng file Excel tổng hợp + dashboard HTML từ 2 file xlsx (ngày/tháng) + thống kê mất sóng."""
import os, json, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
HERE=os.path.dirname(os.path.abspath(__file__))
PROV="Quang Ngai"

def _rows(ws, hr, prov):
    H=[ws.cell(hr,c).value for c in range(1,47)]
    def ci(name):
        for i,h in enumerate(H):
            if h and name==str(h).strip(): return i
        for i,h in enumerate(H):
            if h and name in str(h): return i
        return None
    return H, ci

def extract_kpi(day_xlsx, month_xlsx):
    wb=openpyxl.load_workbook(day_xlsx, data_only=True); ws=wb.active; hr=3
    H,ci=_rows(ws,hr,PROV)
    keys={'xa':'Tỷ lệ Xã/phường đạt MTCL 2026','mtcl':'MTCL_2026','dg':'Đánh giá','kpi':'HTMT_KPI',
          'qos':'HTMT_QOS','vhkt':'HTMT_VHKT','mll':'MLL_TIME','traffic':'Traffic All (GB)','sotram':'Số lượng trạm'}
    ix={k:ci(v) for k,v in keys.items()}
    daily=[]
    for r in range(hr+1, ws.max_row+1):
        prov=ws.cell(r,4).value
        if prov and PROV in str(prov):
            v=[ws.cell(r,c).value for c in range(1,47)]; s=str(v[1])
            daily.append({'date':'%s/%s'%(s[6:8],s[4:6]),'xa':v[ix['xa']],'mtcl':v[ix['mtcl']],'dg':v[ix['dg']],
                'kpi':v[ix['kpi']],'qos':v[ix['qos']],'vhkt':v[ix['vhkt']],'mll':v[ix['mll']],'traffic':v[ix['traffic']]})
    daily.sort(key=lambda x:x['date'])
    wm=openpyxl.load_workbook(month_xlsx, data_only=True); wsm=wm.active; mr=4
    HM,cim=_rows(wsm,mr,PROV)
    mk={'xa':'Tỷ lệ Xã','mtcl':'MTCL_2026','dg':'Đánh giá','kpi':'HTMT_KPI','qos':'HTMT_QOS',
        'vhkt':'HTMT_VHKT','mll':'MLL_TIME','sotram':'Số lượng trạm','sl_xa':'SL Xã'}
    mix={k:cim(v) for k,v in mk.items()}
    monthly=None
    for r in range(mr+1, wsm.max_row+1):
        if wsm.cell(r,4).value and PROV in str(wsm.cell(r,4).value):
            v=[wsm.cell(r,c).value for c in range(1,47)]
            monthly={k:v[mix[k]] for k in mk}; break
    return daily, monthly

def build_excel(day_xlsx, month_xlsx, ms, out_path):
    thin=Side('thin', color='BFBFBF'); bd=Border(thin,thin,thin,thin)
    hf=PatternFill('solid', fgColor='1F4E79'); hff=Font(color='FFFFFF', bold=True, size=10)
    tf=Font(bold=True, size=12, color='1F4E79'); cf=Font(size=10)
    cen=Alignment('center','center',wrap_text=True); lef=Alignment('left','center')
    out=openpyxl.Workbook()
    def raw_sheet(ws_src, hr, title, name, wb):
        ws=wb.create_sheet(name) if name not in wb.sheetnames else wb[name]
        H=[ws_src.cell(hr,c).value for c in range(1,47)]
        cols=[i for i,h in enumerate(H) if h is not None]
        ws.append([title]); ws['A1'].font=tf; ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
        ws.append([H[i] for i in cols])
        for j in range(1,len(cols)+1):
            c=ws.cell(2,j); c.fill=hf; c.font=hff; c.alignment=cen; c.border=bd
        for r in range(hr+1, ws_src.max_row+1):
            prov=ws_src.cell(r,4).value
            if prov and PROV in str(prov):
                ws.append([ws_src.cell(r,i+1).value for i in cols]); rr=ws.max_row
                for j in range(1,len(cols)+1):
                    cc=ws.cell(rr,j); cc.font=cf; cc.border=bd; cc.alignment=lef if j<=4 else cen
        for j,i in enumerate(cols,1):
            ws.column_dimensions[get_column_letter(j)].width=min(max(len(str(H[i]))+2,9),22)
        ws.freeze_panes=ws.cell(3,1)
    wbd=openpyxl.load_workbook(day_xlsx, data_only=True)
    ws1=out.active; ws1.title='KPI_Ngay_QuangNgai'
    raw_sheet(wbd.active, 3, 'MTCL2026 - Quảng Ngãi | KPI theo ngày', 'KPI_Ngay_QuangNgai', out)
    wbm=openpyxl.load_workbook(month_xlsx, data_only=True)
    raw_sheet(wbm.active, 4, 'MTCL2026 - Quảng Ngãi | Lũy kế tháng', 'LuyKe_Thang_QuangNgai', out)
    # MS sheets
    def ms_sheet(name, title, cols, rows):
        ws=out.create_sheet(name)
        ws.append([title]); ws['A1'].font=Font(bold=True,size=12,color='B71C1C'); ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=len(cols))
        ws.append(cols)
        for j in range(1,len(cols)+1):
            c=ws.cell(2,j); c.fill=PatternFill('solid',fgColor='B71C1C'); c.font=hff; c.alignment=cen; c.border=bd
        for row in rows:
            ws.append(row); r=ws.max_row
            for j in range(1,len(cols)+1):
                c=ws.cell(r,j); c.font=cf; c.border=bd; c.alignment=lef if j<=3 else cen
        for j,cn in enumerate(cols,1):
            ws.column_dimensions[get_column_letter(j)].width=min(max(len(str(cn))+2,10),34)
        ws.freeze_panes=ws.cell(3,1)
    ms_sheet('MatSong_TramKeoDai','Top trạm mất liên lạc kéo dài',
        ['Site','Lớp','Khu vực','Số phút','Trạng thái','Bắt đầu','Nguyên nhân'],
        [[x['site'],x['net'],x['district'],x['min'],'Đang mất' if x['ongoing'] else 'Đã KP',x['start'],x['cause'] or '-'] for x in ms['longest']])
    ms_sheet('MatSong_TanSuat','Top trạm mất sóng nhiều lần',
        ['Site','Khu vực','Lớp mạng','Số lần','Tổng phút','Dài nhất'],
        [[x['site'],x['district'],x['nets'],x['count'],x['total'],x['max']] for x in ms['freq']])
    ms_sheet('MatSong_DangMat','Trạm đang mất liên lạc',
        ['Site','Lớp','Khu vực','Số phút','Bắt đầu','Nguyên nhân'],
        [[x['site'],x['net'],x['district'],x['min'],x['start'],x['cause'] or '-'] for x in ms['ongoing_list']])
    out.save(out_path)

def build_dashboard(daily, monthly, ms, out_path, gen_ts=None):
    import datetime
    if gen_ts is None:
        gen_ts=datetime.datetime.now().strftime('%d/%m/%Y %H:%M')
    DATA=json.dumps({'daily':daily,'monthly':monthly,'ms':ms}, ensure_ascii=False)
    tpl=open(os.path.join(HERE,'dashboard_template.html'), encoding='utf-8').read()
    js=open(os.path.join(HERE,'app_template.js'), encoding='utf-8').read()
    # Nội suy Chart.js để file tự chứa, chia sẻ được cả khi không có mạng
    chartjs_path=os.path.join(HERE,'vendor','chart.umd.min.js')
    chartjs=open(chartjs_path, encoding='utf-8').read() if os.path.exists(chartjs_path) else ''
    html=(tpl.replace('__CHARTJS__', chartjs)
             .replace('__GENTS__', gen_ts)
             .replace('__DATA__', DATA)
             .replace('<script src="app.js"></script>', '<script>\n'+js+'\n</script>'))
    open(out_path,'w',encoding='utf-8').write(html)
